import io
import pytest
from fastapi import HTTPException
try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None
import docx  # python-docx

from app.services.extractor import (
    validate_file,
    extract_text,
    extract_text_from_pdf,
    extract_text_from_docx,
    extract_text_from_txt,
)


def test_validate_file_valid():
    assert validate_file("test.pdf", 1000) == ".pdf"
    assert validate_file("notes.TXT", 500) == ".txt"
    assert validate_file("report.docx", 2048) == ".docx"


def test_validate_file_unsupported_extension():
    with pytest.raises(HTTPException) as exc_info:
        validate_file("malicious.exe", 1000)
    assert exc_info.value.status_code == 400
    assert "Unsupported file format" in exc_info.value.detail


def test_validate_file_empty():
    with pytest.raises(HTTPException) as exc_info:
        validate_file("empty.txt", 0)
    assert exc_info.value.status_code == 400
    assert "empty" in exc_info.value.detail


def test_validate_file_oversized():
    with pytest.raises(HTTPException) as exc_info:
        validate_file("huge.pdf", 11 * 1024 * 1024)
    assert exc_info.value.status_code == 413
    assert "exceeds maximum allowed size" in exc_info.value.detail


def test_extract_text_from_txt():
    content = "This is a comprehensive test document designed to verify plain text extraction.".encode("utf-8")
    result = extract_text_from_txt(content)
    assert "comprehensive test document" in result


def test_extract_text_from_txt_latin1():
    content = "Café au lait and résumé extraction in latin-1 format.".encode("latin-1")
    result = extract_text_from_txt(content)
    assert "Café" in result or "extraction" in result


@pytest.mark.skipif(fitz is None, reason="PyMuPDF not installed")
def test_extract_text_from_pdf_digital():
    # Create in-memory PDF using PyMuPDF
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Hello from intelligent document summarizer on Amazon SageMaker AI!")
    pdf_bytes = doc.tobytes()
    doc.close()

    result = extract_text_from_pdf(pdf_bytes)
    assert "intelligent document summarizer" in result


def test_extract_text_from_pdf_scanned_rejection():
    # Create blank PDF without text (simulating image-only or empty scanned document)
    import pypdf
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=300, height=300)
    bio = io.BytesIO()
    writer.write(bio)
    pdf_bytes = bio.getvalue()

    with pytest.raises(HTTPException) as exc_info:
        extract_text_from_pdf(pdf_bytes)
    assert exc_info.value.status_code == 422
    assert "OCR" in exc_info.value.detail


def test_extract_text_from_docx():
    # Create in-memory DOCX
    doc = docx.Document()
    doc.add_paragraph("This is the main introduction paragraph of the document.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Value B"
    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    result = extract_text_from_docx(docx_bytes)
    assert "main introduction paragraph" in result
    assert "Header A | Value B" in result


def test_extract_text_pipeline():
    sample_text = (
        "Cloud computing is the on-demand delivery of IT resources over the Internet with pay-as-you-go pricing. "
        "Instead of buying, owning, and maintaining physical data centers and servers, you can access technology services, "
        "such as computing power, storage, and databases, on an as-needed basis from a cloud provider like Amazon Web Services."
    )
    res = extract_text("cloud_intro.txt", sample_text.encode("utf-8"))
    assert res.word_count > 20
    assert res.char_count == len(sample_text)
    assert res.truncated is False
